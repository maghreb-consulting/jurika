# -*- coding: utf-8 -*-
"""
Phase C — Intégration des 3 modèles directeur Dissolution / Liquidation :
  - PV_DISSOLUTION_LIQUIDATION_SARL      (src Downloads/3_..._SARL_modele_deterministe.docx)
  - PV_DISSOLUTION_LIQUIDATION_SARL_AU   (src Downloads/4_..._SARL_AU_modele_deterministe.docx)
  - RAPPORT_LIQUIDATION_DIRECTEUR        (src Desktop/model/RAPPORT_LIQUIDATION_SARL_et_SARL_AU_modele_deterministe.docx)

Traitement identique aux Phases 4 / A / B (script réutilisable) :
  1. Retrait du PRÉAMBULE méta : tout ce qui précède le 1er paragraphe « $DENOMINATION »
     (en-tête « MODÈLE DÉTERMINISTE », Réf/Workflow/Objet/Base, CONVENTION DE BALISAGE,
     légende, règle d'uniformité, titre méta « EN-TÊTE »).
  2. Retrait de l'ANNEXE glossaire : paragraphe « ANNEXE — … » / « DICTIONNAIRE DES
     VARIABLES … » et tout ce qui suit (sectPr final conservé). Le RAPPORT n'a pas d'annexe
     de fin : rien n'est retiré côté queue dans ce cas.
  3. Restyle charte JURIKA (Times, en-tête société centré navy, titres de section
     JurikaTitreArticle, corps justifié, marges 2,5 cm, pied numéroté) — le contenu
     juridique directeur n'est JAMAIS modifié. Réexécutable (recopie depuis la source).
"""
import os
import io
import sys
import copy
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HOME = os.path.expanduser("~")
DOWNLOADS = os.path.join(HOME, "Downloads")
DESKTOP_MODEL = os.path.join(HOME, "Desktop", "model")
DOCX_DIR = os.path.abspath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "backend-java", "ai-service", "src", "main", "resources", "templates", "docx",
))

# Style shell (Times docDefaults + JurikaTitreArticle) transplanté depuis un PV Phase 4.
SHELL_PV = os.path.join(DOCX_DIR, "PV_DEFAUT_QUORUM_SARL.docx")

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x66, 0x66, 0x66)

SOCIETE_PREFIXES = (
    "SOCIÉTÉ À RESPONSABILITÉ LIMITÉE",
    "AU CAPITAL DE",
    "SIÈGE SOCIAL",
    "IMMATRICULÉE AU REGISTRE",
    "RC N°",
    "SOCIÉTÉ EN LIQUIDATION",
)
DOCTITLE_PREFIXES = (
    "PROCÈS-VERBAL",
    "RAPPORT DE LIQUIDATION",
)
CONTROL_PREFIXES = ("◇", "◆", "▼", "▲", "▶", "◁")
STRIP_ANNEX_PREFIXES = ("DICTIONNAIRE DES VARIABLES", "ANNEXE")


def para_text(p):
    el = getattr(p, "_p", p)
    return "".join(n.text or "" for n in el.iter(qn("w:t")))


def source_style_id(p):
    ppr = p._p.find(qn("w:pPr"))
    if ppr is None:
        return None
    ps = ppr.find(qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


# ----------------------------------------------------------------------------
# 1 + 2 : strip préambule / annexe glossaire (niveau body, préserve sectPr)
# ----------------------------------------------------------------------------
def strip_preamble_and_annex(doc):
    body = doc.element.body
    children = list(body.iterchildren())

    start = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p") and para_text(ch).strip().startswith("$DENOMINATION"):
            start = i
            break
    if start is None:
        raise RuntimeError("Marqueur $DENOMINATION introuvable — préambule non retiré.")

    annex = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p"):
            t = para_text(ch).strip()
            if any(t.startswith(pref) for pref in STRIP_ANNEX_PREFIXES):
                annex = i
                break

    sectpr = None
    if children and children[-1].tag == qn("w:sectPr"):
        sectpr = children[-1]

    end = annex if annex is not None else len(children)
    keep = children[start:end]

    for ch in children:
        body.remove(ch)
    for ch in keep:
        body.append(ch)
    if sectpr is not None:
        body.append(sectpr)


# ----------------------------------------------------------------------------
# 3 : styles.xml — docDefaults Times + style JurikaTitreArticle (transplanté)
# ----------------------------------------------------------------------------
def ensure_doc_defaults_times(doc):
    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        dd = OxmlElement("w:docDefaults")
        styles_el.insert(0, dd)
    rprd = dd.find(qn("w:rPrDefault"))
    if rprd is None:
        rprd = OxmlElement("w:rPrDefault")
        dd.insert(0, rprd)
    rpr = rprd.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rprd.append(rpr)
    el = rpr.find(qn("w:rFonts"))
    if el is not None:
        rpr.remove(el)
    rfonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    rpr.insert(0, rfonts)
    for tag, val in (("w:sz", "22"), ("w:szCs", "22")):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), val)


def ensure_jurika_titre_style(doc):
    styles_el = doc.styles.element
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "JurikaTitreArticle":
            return
    shell = Document(SHELL_PV)
    src = None
    for st in shell.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "JurikaTitreArticle":
            src = st
            break
    if src is None:
        raise RuntimeError("JurikaTitreArticle absent du shell PV.")
    styles_el.append(copy.deepcopy(src))


# ----------------------------------------------------------------------------
# 3 : restyle runs + paragraphes
# ----------------------------------------------------------------------------
def set_run_font(run, color):
    run.font.name = "Times New Roman"
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    if color is not None:
        run.font.color.rgb = color


def set_pstyle(p, style_id):
    ppr = p._p.get_or_add_pPr()
    ps = ppr.find(qn("w:pStyle"))
    if ps is None:
        ps = OxmlElement("w:pStyle")
        ppr.insert(0, ps)
    ps.set(qn("w:val"), style_id)


def classify(p):
    if source_style_id(p) == "Heading1":
        return "section"
    t = para_text(p).strip()
    if not t:
        return "body"
    if t.startswith(CONTROL_PREFIXES):
        return "control"
    if t.startswith("$DENOMINATION"):
        return "denom"
    if any(t.startswith(pref) for pref in DOCTITLE_PREFIXES):
        return "doctitle"
    if any(t.startswith(pref) for pref in SOCIETE_PREFIXES):
        return "societe"
    return "body"


def restyle_paragraph(p):
    kind = classify(p)
    if kind == "denom":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            set_run_font(r, NAVY)
    elif kind == "societe":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = False
            set_run_font(r, BLACK)
    elif kind == "doctitle":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            set_run_font(r, BLACK)
    elif kind == "section":
        set_pstyle(p, "JurikaTitreArticle")
        for r in p.runs:
            set_run_font(r, NAVY)
    elif kind == "control":
        for r in p.runs:
            set_run_font(r, BLACK)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            set_run_font(r, BLACK)


def restyle_all(doc):
    for p in doc.paragraphs:
        restyle_paragraph(p)


# ----------------------------------------------------------------------------
# 3 : marges + pied de page numéroté
# ----------------------------------------------------------------------------
def set_margins_and_footer(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)
        footer = s.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        set_run_font(run, GREY)
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(instr); run._r.append(f2)


# ----------------------------------------------------------------------------
# Orchestration
# ----------------------------------------------------------------------------
def integrate(src_path, dest_code, must_contain):
    dest = os.path.join(DOCX_DIR, dest_code + ".docx")
    if not os.path.exists(src_path):
        raise RuntimeError("Source introuvable : " + src_path)
    shutil.copyfile(src_path, dest)
    doc = Document(dest)
    strip_preamble_and_annex(doc)
    ensure_doc_defaults_times(doc)
    ensure_jurika_titre_style(doc)
    restyle_all(doc)
    set_margins_and_footer(doc)
    doc.save(dest)

    check = Document(dest)
    full = "\n".join(para_text(p) for p in check.paragraphs)
    assert "CONVENTION DE BALISAGE" not in full, "préambule résiduel dans " + dest_code
    assert "ANNEXE — VARIABLES" not in full, "annexe glossaire résiduelle dans " + dest_code
    assert "MODÈLE DÉTERMINISTE" not in full, "en-tête méta résiduel dans " + dest_code
    assert "$DENOMINATION" in full, "marqueur $DENOMINATION perdu dans " + dest_code
    for v in must_contain:
        assert v in full, "variable %s perdue dans %s" % (v, dest_code)
    print("  OK -> %s (%d paragraphes)" % (dest_code, len(check.paragraphs)))


def main():
    print("Intégration modèles Dissolution / Liquidation dans:", DOCX_DIR)
    integrate(
        os.path.join(DOWNLOADS, "3_PV_DISSOLUTION_LIQUIDATION_SARL_modele_deterministe.docx"),
        "PV_DISSOLUTION_LIQUIDATION_SARL",
        ("$PV_ETAPE", "$DISSOLUTION_MOTIF", "$LIQUIDATEUR_NOM", "$LIQ_RESULTAT_SENS"),
    )
    integrate(
        os.path.join(DOWNLOADS, "4_PV_DISSOLUTION_LIQUIDATION_SARL_AU_modele_deterministe.docx"),
        "PV_DISSOLUTION_LIQUIDATION_SARL_AU",
        ("$PV_ETAPE", "$DISSOLUTION_MOTIF", "$LIQUIDATEUR_NOM", "$LIQ_RESULTAT_SENS"),
    )
    integrate(
        os.path.join(DESKTOP_MODEL, "RAPPORT_LIQUIDATION_SARL_et_SARL_AU_modele_deterministe.docx"),
        "RAPPORT_LIQUIDATION_DIRECTEUR",
        ("$ACTIF_REALISE_CHIFFRES", "$PASSIF_REGLE_CHIFFRES", "$BONI_LIQUIDATION_CHIFFRES",
         "$MALI_LIQUIDATION_CHIFFRES", "$LIQUIDATEUR_GENRE", "$ASSOCIE_BONI_CHIFFRES"),
    )
    print("Terminé.")


if __name__ == "__main__":
    main()
