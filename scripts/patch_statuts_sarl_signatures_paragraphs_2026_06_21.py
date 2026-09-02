# -*- coding: utf-8 -*-
"""Sprint Cowork 2026-06-21 (F — signatures simples) — Remplace le TABLEAU
signature par des PARAGRAPHES repetables.

Structure injectee (apres "SIGNATURES") :
    ▶ ASSOCIES_SIGNATURE {{signature_nom}}
    ▶ ASSOCIES_SIGNATURE Signature : ____________________
    ▶ ASSOCIES_SIGNATURE
    ▶ GERANTS_SIGNATURE {{signature_nom}}
    ▶ GERANTS_SIGNATURE « Bon pour acceptation des fonctions »
    ▶ GERANTS_SIGNATURE Signature : ____________________
    ▶ GERANTS_SIGNATURE

L'engine (expandBodyParagraphs Option A) detecte les paragraphes contigus
portant le MEME ▶ NAME et les clone par item. Chaque associe -> 3 paragraphes
clones ; chaque gerant -> 4 paragraphes clones (avec la mention en plus).
Le 'paragraphe vide' final donne un saut de ligne entre 2 signatures.

Usage :
    PYTHONIOENCODING=utf-8 python scripts/patch_statuts_sarl_signatures_paragraphs_2026_06_21.py
"""
import io
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
TPL_DIR = ROOT / "backend-java" / "ai-service" / "src" / "main" / "resources" / "templates" / "docx"


def remove_signature_table(doc) -> bool:
    """Supprime la 1ere table contenant le marqueur ▶ ASSOCIES_SIGNATURE."""
    body = doc.element.body
    for tbl in list(doc.tables):
        if "ASSOCIES_SIGNATURE" in tbl._tbl.xml:
            body.remove(tbl._tbl)
            return True
    return False


def find_signatures_heading(doc):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "SIGNATURES":
            return i, p
    return -1, None


def add_paragraph_after(par, text: str):
    """Insere un paragraphe APRES `par` en heritant du style."""
    import copy
    new_p = copy.deepcopy(par._p)
    # Vide tous les runs
    for r in new_p.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"
    ):
        new_p.remove(r)
    par._p.addnext(new_p)
    new_par = par.__class__(new_p, par._parent)
    new_par.add_run(text)
    return new_par


def patch(path: Path) -> None:
    print(f"\n=== {path.name} ===")
    doc = Document(str(path))
    if remove_signature_table(doc):
        print("  [OK] tableau signature supprime")
    else:
        print("  [SKIP] pas de tableau signature trouve")
    idx, par = find_signatures_heading(doc)
    if par is None:
        print("  [SKIP] heading 'SIGNATURES' introuvable")
        doc.save(str(path))
        return
    if "ASSOCIES_SIGNATURE" in '\n'.join(p.text for p in doc.paragraphs[idx:idx+15]):
        print("  [SKIP] paragraphes signature deja injectes")
        doc.save(str(path))
        return

    # Insertion DANS L'ORDRE -> chaque add_paragraph_after pose juste apres `par`,
    # donc on ajoute en SENS INVERSE pour que l'ordre final soit ▶ASSO×3 puis ▶GER×4.
    lines = [
        # Gerants (en dernier visuellement -> insere en premier ici)
        "▶ GERANTS_SIGNATURE ",
        "▶ GERANTS_SIGNATURE Signature : ____________________",
        "▶ GERANTS_SIGNATURE « Bon pour acceptation des fonctions »",
        "▶ GERANTS_SIGNATURE {{signature_nom}}",
        # Associes
        "▶ ASSOCIES_SIGNATURE ",
        "▶ ASSOCIES_SIGNATURE Signature : ____________________",
        "▶ ASSOCIES_SIGNATURE {{signature_nom}}",
    ]
    for line in lines:
        add_paragraph_after(par, line)
    print(f"  [OK] {len(lines)} paragraphes signature injectes apres SIGNATURES")
    doc.save(str(path))
    print(f"  [OK] sauvegarde {path.name}")


for fn in ["STATUTS_CONSTITUTIFS_SARL.docx", "STATUTS_CONSTITUTIFS_SARL_AU.docx"]:
    patch(TPL_DIR / fn)
