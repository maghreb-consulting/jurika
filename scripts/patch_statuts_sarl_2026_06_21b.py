# -*- coding: utf-8 -*-
"""Sprint Cowork 2026-06-21 (E — contenu statuts) — Patches Article 29
ACOMPTES_DIVIDENDES (toujours visible) + remplacement du texte signatures
par un tableau cabinet nominatif (▶ ASSOCIES_SIGNATURE + ▶ GERANTS_SIGNATURE).

Cibles :
  - STATUTS_CONSTITUTIFS_SARL.docx     (Article 29 + signatures)
  - STATUTS_CONSTITUTIFS_SARL_AU.docx  (Article 26 + signature unique)

Usage :
    PYTHONIOENCODING=utf-8 python scripts/patch_statuts_sarl_2026_06_21b.py
"""
import copy
import io
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document  # type: ignore
from docx.oxml.ns import qn  # type: ignore
from lxml import etree  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
TPL_DIR = ROOT / "backend-java" / "ai-service" / "src" / "main" / "resources" / "templates" / "docx"


def set_paragraph_text(par, new_text: str) -> None:
    """Remplace le texte d'un paragraphe en preservant le style du 1er run."""
    if not par.runs:
        par.add_run(new_text)
        return
    first = par.runs[0]
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)
    first.text = new_text


def find_para(doc, predicate, start: int = 0):
    for i in range(start, len(doc.paragraphs)):
        if predicate(doc.paragraphs[i].text):
            return i, doc.paragraphs[i]
    return -1, None


def strip_markers(text: str, *markers: str) -> str:
    out = text
    for m in markers:
        out = out.replace(m, "")
    while "  " in out:
        out = out.replace("  ", " ")
    return out.strip()


def patch_acomptes(doc, label: str) -> None:
    """Article 29/26 — strip ◇/◆ ACOMPTES_DIVIDENDES pour TOUJOURS afficher
    le corps legal (acomptes = sous-option, le reste est obligatoire)."""
    for i, p in enumerate(doc.paragraphs):
        if "ACOMPTES_DIVIDENDES" in p.text:
            new = strip_markers(p.text, "◇ ACOMPTES_DIVIDENDES", "◆ ACOMPTES_DIVIDENDES")
            set_paragraph_text(p, new)
            print(f"[OK] {label} para {i} : markers ACOMPTES_DIVIDENDES stripped")
            return
    print(f"[SKIP] {label} : aucun marker ACOMPTES_DIVIDENDES (deja fait ?)")


def add_signature_table(doc, signature_text_predicate, label: str) -> None:
    """Remplace le paragraphe de signature par un tableau cabinet 3 colonnes
    (Qualite / Nom / Mention + Signature) avec 2 marqueurs ▶ ASSOCIES_SIGNATURE
    et ▶ GERANTS_SIGNATURE que le moteur clone."""
    idx, par = find_para(doc, signature_text_predicate)
    if par is None:
        print(f"[SKIP] {label} : paragraphe signature introuvable")
        return
    if "ASSOCIES_SIGNATURE" in par.text:
        print(f"[SKIP] {label} : tableau signature deja injecte")
        return

    # Header explicatif (remplace le texte d'instruction).
    set_paragraph_text(par, "SIGNATURES")

    # Construire le tableau via XML brut (python-docx ne permet pas d'inserer
    # un tableau a une position arbitraire facilement ; on cree via add_table
    # puis on deplace l'element XML juste apres `par`).
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'

    # Ligne 0 = header.
    hdr = table.rows[0].cells
    hdr[0].text = "Qualité"
    hdr[1].text = "Nom et prénom"
    hdr[2].text = "Mention manuscrite + Signature"
    for c in hdr:
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True

    # Ligne 1 = ASSOCIES_SIGNATURE (clonee N fois par le moteur).
    r1 = table.rows[1].cells
    r1[0].text = "▶ ASSOCIES_SIGNATURE {{signature_titre}}"
    r1[1].text = "{{signature_nom}}"
    r1[2].text = "{{signature_mention}}\n\nSignature : ____________________"

    # Ligne 2 = GERANTS_SIGNATURE.
    r2 = table.rows[2].cells
    r2[0].text = "▶ GERANTS_SIGNATURE {{signature_titre}}"
    r2[1].text = "{{signature_nom}}"
    r2[2].text = "{{signature_mention}}\n\nSignature : ____________________"

    # Deplacer la table apres `par` (par defaut add_table l'ajoute en fin de body).
    body = doc.element.body
    tbl_el = table._tbl
    body.remove(tbl_el)
    par._p.addnext(tbl_el)
    print(f"[OK] {label} : tableau signature injecte (3 lignes, 3 colonnes)")


# === SARL ===
SARL = TPL_DIR / "STATUTS_CONSTITUTIFS_SARL.docx"
doc = Document(str(SARL))
patch_acomptes(doc, "SARL Art.29")
add_signature_table(
    doc,
    lambda t: t.strip().startswith("Signatures des associés"),
    "SARL",
)
doc.save(str(SARL))
print(f"[OK] sauvegarde {SARL.name}\n")

# === SARL_AU ===
SARL_AU = TPL_DIR / "STATUTS_CONSTITUTIFS_SARL_AU.docx"
doc = Document(str(SARL_AU))
patch_acomptes(doc, "SARL_AU Art.26")
add_signature_table(
    doc,
    lambda t: t.strip().startswith("Signature de l'associé unique"),
    "SARL_AU",
)
doc.save(str(SARL_AU))
print(f"[OK] sauvegarde {SARL_AU.name}")
