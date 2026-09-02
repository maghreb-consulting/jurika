# -*- coding: utf-8 -*-
"""
Phase A — Intégration des 2 modèles directeur de séance d'AG partagés :
  - CONVOCATION_AG        (source : 5_CONVOCATION_AG_SARL_et_SARL_AU_modele_deterministe.docx)
  - FEUILLE_PRESENCE_AG   (source : 6_FEUILLE_PRESENCE_AG_SARL_et_SARL_AU_modele_deterministe.docx)

Traitement (identique à la Phase 4 — 8 modèles) :
  1. Retrait du PRÉAMBULE (tout ce qui précède le 1er paragraphe « $DENOMINATION » :
     entête méta + CONVENTION DE BALISAGE).
  2. Retrait de l'ANNEXE — VARIABLES NOUVELLES (le paragraphe Heading1 « ANNEXE — … »
     et tout ce qui suit, jusqu'au <w:sectPr> final conservé).
  3. Restyle charte JURIKA : police Times New Roman partout (docDefaults + runs), titres
     navy 1F4E79 / corps noir, entête société centré, style JurikaTitreArticle sur les
     titres, corps justifié, marges 2,5 cm, pied de page avec numérotation.
  4. FEUILLE uniquement : la boucle ASSOCIES est portée par un TABLEAU au format directeur
     (ligne ▼ / ligne données / ligne ▲). Le moteur ne sait dérouler une boucle de TABLE
     qu'au format LEGACY « ▶ NOM » porté par la ligne de données elle-même. On collapse
     donc les 3 lignes en 1 : suppression des lignes marqueur ▼ et ▲, préfixe « ▶ ASSOCIES »
     injecté dans la 1re cellule de la ligne de données. (Contenu directeur inchangé :
     seule la plomberie de balisage de boucle est adaptée à l'intégration.)

Le contenu juridique (prose directeur) n'est JAMAIS modifié. Réexécutable : recopie
systématiquement depuis la source.
"""
import os
import re
import sys
import io
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HOME = os.path.expanduser("~")
SRC_DIR = os.path.join(HOME, "Desktop", "model")
DOCX_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "backend-java", "ai-service", "src", "main", "resources", "templates", "docx",
)
DOCX_DIR = os.path.abspath(DOCX_DIR)

# Style shell (Times docDefaults + JurikaTitreArticle) transplanté depuis un PV Phase 4.
SHELL_PV = os.path.join(DOCX_DIR, "PV_DEFAUT_QUORUM_SARL.docx")

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLACK = RGBColor(0x00, 0x00, 0x00)

# Détection entête société (centré, navy, gras).
HEADER_PREFIXES = (
    "$DENOMINATION",
    "SOCIÉTÉ À RESPONSABILITÉ LIMITÉE",
    "AU CAPITAL DE",
    "SIÈGE SOCIAL",
    "R.C.",
)
# Détection titre principal (JurikaTitreArticle).
TITLE_PREFIXES = (
    "FEUILLE DE PRÉSENCE À L'ASSEMBLÉE",
)
CONTROL_PREFIXES = ("◇", "◆", "▼", "▲", "▶", "◀")


def para_text(p):
    el = getattr(p, "_p", p)  # accepte Paragraph python-docx ou élément lxml <w:p>
    return "".join(n.text or "" for n in el.iter(qn("w:t")))


def p_style(p):
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        return None
    ps = ppr.find(qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


# ----------------------------------------------------------------------------
# 1 + 2 : strip préambule / annexe (niveau body, préserve tables & sectPr)
# ----------------------------------------------------------------------------
def strip_preamble_and_annex(doc):
    body = doc.element.body
    children = list(body.iterchildren())

    # Index du 1er <w:p> dont le texte commence par $DENOMINATION.
    start = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p") and para_text(ch).strip().startswith("$DENOMINATION"):
            start = i
            break
    if start is None:
        raise RuntimeError("Marqueur $DENOMINATION introuvable — préambule non retiré.")

    # Index de l'ANNEXE (paragraphe dont le texte commence par 'ANNEXE').
    annex = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p") and para_text(ch).strip().startswith("ANNEXE"):
            annex = i
            break

    # sectPr final à préserver (dernier enfant si c'est un sectPr).
    sectpr = None
    if children and children[-1].tag == qn("w:sectPr"):
        sectpr = children[-1]

    # Fenêtre à conserver : [start .. annex) puis sectPr.
    end = annex if annex is not None else len(children)
    keep = children[start:end]

    for ch in children:
        body.remove(ch)
    for ch in keep:
        body.append(ch)
    if sectpr is not None:
        body.append(sectpr)


# ----------------------------------------------------------------------------
# 3 : styles.xml — docDefaults Times + style JurikaTitreArticle
# ----------------------------------------------------------------------------
def ensure_doc_defaults_times(doc):
    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        dd = OxmlElement("w:docDefaults")
        styles_el.insert(0, dd)
    rprd = dd.find(qn("w:rPrDefault"))
    if rprd is None:
        rprd = OxmlElement("w:rPrDefault")
        dd.insert(0, rprd)
    rpr = rprd.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rprd.append(rpr)
    for tag in ("w:rFonts",):
        el = rpr.find(qn(tag))
        if el is not None:
            rpr.remove(el)
    rfonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    rpr.insert(0, rfonts)
    for tag, val in (("w:sz", "22"), ("w:szCs", "22")):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), val)


def ensure_jurika_titre_style(doc):
    styles_el = doc.styles.element
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "JurikaTitreArticle":
            return
    # Extrait la définition exacte depuis le shell PV.
    shell = Document(SHELL_PV)
    src = None
    for st in shell.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "JurikaTitreArticle":
            src = st
            break
    if src is None:
        raise RuntimeError("JurikaTitreArticle absent du shell PV.")
    import copy
    styles_el.append(copy.deepcopy(src))


# ----------------------------------------------------------------------------
# 3 : restyle runs + paragraphes
# ----------------------------------------------------------------------------
def set_run_font(run, color):
    run.font.name = "Times New Roman"
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    if color is not None:
        run.font.color.rgb = color


def set_pstyle(p, style_id):
    ppr = p._p.get_or_add_pPr()
    ps = ppr.find(qn("w:pStyle"))
    if ps is None:
        ps = OxmlElement("w:pStyle")
        ppr.insert(0, ps)
    ps.set(qn("w:val"), style_id)


def classify(text):
    t = text.strip()
    if not t:
        return "body"
    if t.startswith(CONTROL_PREFIXES):
        return "control"
    for pref in TITLE_PREFIXES:
        if t.startswith(pref):
            return "title"
    for pref in HEADER_PREFIXES:
        if t.startswith(pref):
            return "header"
    return "body"


def restyle_paragraph(p):
    kind = classify(para_text(p))
    if kind == "header":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            set_run_font(r, NAVY)
    elif kind == "title":
        set_pstyle(p, "JurikaTitreArticle")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            set_run_font(r, NAVY)
    elif kind == "control":
        for r in p.runs:
            set_run_font(r, BLACK)
    else:  # body
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            set_run_font(r, BLACK)


def restyle_all(doc):
    for p in doc.paragraphs:
        restyle_paragraph(p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    # Cellules : Times, noir, pas de justification forcée.
                    for r in p.runs:
                        set_run_font(r, BLACK)


# ----------------------------------------------------------------------------
# 3 : marges + pied de page numéroté
# ----------------------------------------------------------------------------
def set_margins_and_footer(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)
        footer = s.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        # purge d'un éventuel contenu existant
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        set_run_font(run, RGBColor(0x66, 0x66, 0x66))
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(instr); run._r.append(f2)


# ----------------------------------------------------------------------------
# 4 : FEUILLE — collapse boucle table directeur (3 lignes) -> legacy 1 ligne ▶
# ----------------------------------------------------------------------------
DIR_START = re.compile(r"▼\s*D[EÉ]BUT\s+BOUCLE\s*[—–-]\s*([A-Z][A-Z0-9_]*)")
DIR_END = re.compile(r"▲\s*FIN\s+BOUCLE\s*[—–-]\s*([A-Z][A-Z0-9_]*)")


def collapse_table_loop(doc):
    for tbl in doc.tables:
        rows = tbl.rows
        start_i = end_i = None
        name = None
        for i, row in enumerate(rows):
            txt = "\n".join(para_text(p) for c in row.cells for p in c.paragraphs)
            ms = DIR_START.search(txt)
            me = DIR_END.search(txt)
            if ms and start_i is None:
                start_i, name = i, ms.group(1)
            elif me:
                end_i = i
        if start_i is None or end_i is None or name is None:
            continue
        # ligne de données = celle entre start et end (1 seule attendue).
        data_i = start_i + 1
        data_row = tbl.rows[data_i]._tr
        # préfixe ▶ NOM dans la 1re cellule (1er paragraphe).
        first_cell = tbl.rows[data_i].cells[0]
        first_p = first_cell.paragraphs[0]
        marker = OxmlElement("w:r")
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = "▶ %s " % name
        marker.append(t)
        first_p._p.insert(0, marker)
        # supprimer les lignes marqueur ▲ (fin) puis ▼ (début) — ordre décroissant.
        tbl._tbl.remove(tbl.rows[end_i]._tr)
        tbl._tbl.remove(tbl.rows[start_i]._tr)
        print("    [table] boucle %s collapsée (legacy ▶) — lignes ▼/▲ retirées" % name)


# ----------------------------------------------------------------------------
# Orchestration
# ----------------------------------------------------------------------------
def integrate(src_name, dest_code, collapse=False):
    src = os.path.join(SRC_DIR, src_name)
    dest = os.path.join(DOCX_DIR, dest_code + ".docx")
    if not os.path.exists(src):
        raise RuntimeError("Source introuvable : " + src)
    shutil.copyfile(src, dest)
    doc = Document(dest)
    strip_preamble_and_annex(doc)
    ensure_doc_defaults_times(doc)
    ensure_jurika_titre_style(doc)
    if collapse:
        collapse_table_loop(doc)
    restyle_all(doc)
    set_margins_and_footer(doc)
    doc.save(dest)
    # Contrôle rapide : marqueurs préservés, préambule/annexe absents.
    check = Document(dest)
    full = "\n".join(para_text(p) for p in check.paragraphs)
    full += "\n".join(para_text(p) for t in check.tables for row in t.rows
                       for c in row.cells for p in c.paragraphs)
    assert "CONVENTION DE BALISAGE" not in full, "préambule résiduel !"
    assert "ANNEXE" not in full, "annexe résiduelle !"
    assert "$DENOMINATION" in full, "marqueur $ perdu !"
    print("  OK -> %s (%d paragraphes body, %d tables)"
          % (dest_code, len(check.paragraphs), len(check.tables)))


def main():
    print("Intégration modèles séance AG dans:", DOCX_DIR)
    integrate("5_CONVOCATION_AG_SARL_et_SARL_AU_modele_deterministe.docx",
              "CONVOCATION_AG", collapse=False)
    integrate("6_FEUILLE_PRESENCE_AG_SARL_et_SARL_AU_modele_deterministe.docx",
              "FEUILLE_PRESENCE_AG", collapse=True)
    print("Terminé.")


if __name__ == "__main__":
    main()
