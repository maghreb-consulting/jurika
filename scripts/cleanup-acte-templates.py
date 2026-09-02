#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Cleanup Acte templates: remove FORMULAIRE DE SAISIE header + variables table.
Add ▶ ASSOCIES markers for SARL (repeatable blocks). Keep SARL_AU as single-shot.
"""
import io, os, re, sys
from docx import Document
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

D = 'backend-java/ai-service/src/main/resources/templates/docx'


def set_text(p, new_text):
    for r in p.runs:
        r._element.getparent().remove(r._element)
    p.add_run(new_text)


def clean_acte(path, is_au: bool):
    doc = Document(path)
    body = doc.element.body
    # 1. Find the first paragraph that is ${DENOMINATION} - that's the real content start.
    start_idx = None
    children = list(body.iterchildren())
    for i, el in enumerate(children):
        if el.tag.endswith('}p'):
            ts = el.findall('.//' + qn('w:t'))
            joined = ''.join(t.text or '' for t in ts)
            if joined.strip().startswith('${DENOMINATION}'):
                start_idx = i
                break
    if start_idx is None:
        raise SystemExit(f'No ${{DENOMINATION}} in {path}')
    # 2. Remove all body children before start_idx
    sectpr = body.find(qn('w:sectPr'))
    keep = children[start_idx:]
    for el in list(body.iterchildren()):
        body.remove(el)
    for el in keep:
        body.append(el)
    if sectpr is not None and body.find(qn('w:sectPr')) is None:
        body.append(sectpr)
    doc.save(path)

    # 3. Re-open and process: bloc markers
    doc = Document(path)
    # Find "▶ Bloc répétable « ASSOCIES »" lines and the next ${ASSOCIE_NOM} line
    paragraphs = doc.paragraphs
    # We'll mark for deletion the explanatory lines and prepend ▶ ASSOCIES to the next data line.
    to_delete = []
    for i, p in enumerate(paragraphs):
        t = p.text
        # SARL: convert
        if '▶ Bloc répétable' in t or '▶ Signature de' in t or '▶ SARL à associé' in t:
            to_delete.append(p)
            # For SARL only (not AU), find the next data paragraph and prepend ▶ ASSOCIES
            if not is_au:
                # Find next paragraph that contains ${ASSOCIE_NOM} or ${ASSOCIE_PARTS_CHIFFRES}
                for j in range(i + 1, min(i + 4, len(paragraphs))):
                    nt = paragraphs[j].text
                    if '${ASSOCIE_NOM}' in nt or '${ASSOCIE_PARTS_CHIFFRES}' in nt:
                        # Prepend ▶ ASSOCIES if not already there
                        if '▶ ASSOCIES' not in nt:
                            set_text(paragraphs[j], '▶ ASSOCIES ' + nt)
                        break
    for p in to_delete:
        el = p._element
        el.getparent().remove(el)

    doc.save(path)


clean_acte(os.path.join(D, 'ACTE_NOMINATION_GERANT_SARL.docx'), is_au=False)
clean_acte(os.path.join(D, 'ACTE_NOMINATION_GERANT_SARL_AU.docx'), is_au=True)
print('OK')

# Verify
for f in ['ACTE_NOMINATION_GERANT_SARL.docx', 'ACTE_NOMINATION_GERANT_SARL_AU.docx']:
    d = Document(os.path.join(D, f))
    print(f, ':', len(d.paragraphs), 'paragraphs')
    for i, p in enumerate(d.paragraphs[:6]):
        print(f'  {i}: {p.text[:120]}')
    print('  ...')
    for i, p in enumerate(d.paragraphs[14:22]):
        print(f'  {i+14}: {p.text[:120]}')
