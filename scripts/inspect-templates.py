#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Inspect template placeholders and first paragraphs."""
import io, os, re, sys
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

D = 'backend-java/ai-service/src/main/resources/templates/docx'
RE_DOLLAR = re.compile(r'\$\{([^}\n]+?)\}')
RE_BRACE = re.compile(r'\{\{([^}]+?)\}\}')

def collect(doc):
    dollar = set()
    brace = set()
    def scan(p):
        t = p.text
        for m in RE_DOLLAR.findall(t):
            dollar.add(m.strip())
        for m in RE_BRACE.findall(t):
            brace.add(m.strip())
    for p in doc.paragraphs:
        scan(p)
    for tab in doc.tables:
        for row in tab.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    scan(p)
    return sorted(dollar), sorted(brace)

for f in sys.argv[1:]:
    doc = Document(os.path.join(D, f))
    print('='*70)
    print(f'{f} - {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables')
    d, b = collect(doc)
    print(f'  ${{...}} placeholders ({len(d)}):')
    for x in d: print(f'    - {x}')
    print(f'  {{{{...}}}} placeholders ({len(b)}):')
    for x in b: print(f'    - {x}')
    print('  --- first paragraphs ---')
    for i, p in enumerate(doc.paragraphs[:40]):
        if p.text.strip():
            print(f'  P{i}: {p.text[:160]}')
    print('  --- tables ---')
    for ti, tab in enumerate(doc.tables):
        print(f'  TABLE {ti} ({len(tab.rows)} rows x {len(tab.columns)} cols):')
        for ri, row in enumerate(tab.rows[:8]):
            cells_t = ' | '.join(c.text.replace('\n', ' / ')[:60] for c in row.cells)
            print(f'    row{ri}: {cells_t}')
