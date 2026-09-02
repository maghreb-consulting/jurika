# -*- coding: utf-8 -*-
"""Sprint Cowork 2026-06-21 (D — couverture combinatoire).

Patches STATUTS_CONSTITUTIFS_SARL.docx :
  1) Article 12 — wrap TITRE + INTRO en ◇ HAS_CLAUSES_SPECIALES (auto-fermant
     ligne par ligne pour ne pas casser les ◇ CLAUSE_* nichés).
  2) Article 24 — wrap TITRE en ◇ CAC_NOMME (orphan : titre survit alors que
     le corps disparaît si CAC non nommé).
  3) Article 15.2 — remplace le paragraphe single-gérant par :
        - une intro "Sont nommés pour la première fois ..."
        - un bloc répétable ▶ GERANTS ... (un clone par gérant).

Backup : STATUTS_CONSTITUTIFS_SARL.preconv-2026-06-21.docx déjà conservé en
working tree.

Usage :
    PYTHONIOENCODING=utf-8 python scripts/patch_statuts_sarl_2026_06_21.py
"""
import copy
import io
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "backend-java" / "ai-service" / "src" / "main" / "resources" \
    / "templates" / "docx" / "STATUTS_CONSTITUTIFS_SARL.docx"

assert DOCX.exists(), DOCX

doc = Document(str(DOCX))


def set_paragraph_text(par, new_text: str) -> None:
    """Remplace le texte du paragraphe en preservant le style du 1er run.

    Strategy minimale : on garde le 1er run (avec son rPr), on lui injecte
    le nouveau texte, on supprime les autres runs. Les ${...} / {{...}}
    qui auraient ete fragmentes par Word sur plusieurs runs sont ainsi
    re-unifies dans un seul run, ce qui aide aussi la regex du moteur.
    """
    if not par.runs:
        par.add_run(new_text)
        return
    first = par.runs[0]
    # Vider le texte des runs au-dela du 1er
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)
    first.text = new_text


def insert_paragraph_after(par, text: str):
    """Insere un nouveau paragraphe APRES `par` en clonant son CTP pour heriter
    du style (pPr), puis remplace son texte."""
    new_ctp = copy.deepcopy(par._p)
    # Supprimer tous les runs clones (mais garder pPr).
    for r in new_ctp.findall(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
        new_ctp.remove(r)
    par._p.addnext(new_ctp)
    # Re-wrapper en XWPFParagraph python equivalent.
    from docx.oxml.ns import qn
    new_par = par.__class__(new_ctp, par._parent)
    new_par.add_run(text)
    return new_par


def find_para(predicate):
    for i, p in enumerate(doc.paragraphs):
        if predicate(p.text):
            return i, p
    raise AssertionError(f"Paragraphe introuvable: {predicate}")


# --- (1) Article 12 — wrap titre + intro ---
i12_title, p12_title = find_para(
    lambda t: t.strip().startswith("ARTICLE 12 — CLAUSES SPÉCIALES")
)
i12_intro, p12_intro = find_para(
    lambda t: t.strip().startswith("(Activer les clauses ci-après")
)
old = p12_title.text
if "HAS_CLAUSES_SPECIALES" not in old:
    set_paragraph_text(
        p12_title,
        f"◇ HAS_CLAUSES_SPECIALES {old} ◆ HAS_CLAUSES_SPECIALES",
    )
    print(f"[OK] Para {i12_title} (Article 12 titre) wrappe")
else:
    print(f"[SKIP] Para {i12_title} deja wrappe")

old = p12_intro.text
if "HAS_CLAUSES_SPECIALES" not in old:
    set_paragraph_text(
        p12_intro,
        f"◇ HAS_CLAUSES_SPECIALES {old} ◆ HAS_CLAUSES_SPECIALES",
    )
    print(f"[OK] Para {i12_intro} (Article 12 intro) wrappe")
else:
    print(f"[SKIP] Para {i12_intro} deja wrappe")


# --- (2) Article 24 — wrap titre ---
i24, p24 = find_para(
    lambda t: t.strip().startswith("ARTICLE 24 — COMMISSAIRE AUX COMPTES")
)
old = p24.text
if "CAC_NOMME" not in old:
    set_paragraph_text(p24, f"◇ CAC_NOMME {old} ◆ CAC_NOMME")
    print(f"[OK] Para {i24} (Article 24 titre) wrappe")
else:
    print(f"[SKIP] Para {i24} deja wrappe")


# --- (3) Article 15.2 — single-gerant -> intro + bloc ▶ GERANTS ---
i152, p152 = find_para(lambda t: t.strip().startswith("15.2 Premier gérant"))
if "▶ GERANTS" not in p152.text:
    intro = ("15.2 Premiers gérants — Sont nommés pour la première fois "
             "(selon {{gerance_type}}) :")
    set_paragraph_text(p152, intro)
    block_text = (
        "▶ GERANTS {{gerant_civilite}} {{gerant_nom_prenom}}, "
        "de nationalité {{gerant_nationalite}}, "
        "{{gerant_ne}} le {{gerant_date_naissance}} à "
        "{{gerant_lieu_naissance}}, demeurant à {{gerant_adresse}}, "
        "titulaire de la {{gerant_piece_type}} n° {{gerant_piece_numero}}, "
        "pour une durée {{gerant_duree_mandat}} "
        "({{gerant_duree_annees}} années si limitée). Le gérant ainsi nommé "
        "déclare accepter ses fonctions et n'être frappé d'aucune "
        "incapacité ni interdiction. ◀ GERANTS"
    )
    insert_paragraph_after(p152, block_text)
    print(f"[OK] Para {i152} (15.2) splitte en intro + bloc ▶ GERANTS")
else:
    print(f"[SKIP] Para {i152} deja en bloc ▶ GERANTS")


doc.save(str(DOCX))
print(f"\n[OK] Sauvegarde : {DOCX}")
