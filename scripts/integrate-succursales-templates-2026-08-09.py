# -*- coding: utf-8 -*-
"""
Phase D — Intégration des 6 modèles directeur Succursales (création MA / étrangère + fermeture) :
  - PV_CREATION_SUCCURSALE_MAROC_SARL / _AU
  - PV_CREATION_SUCCURSALE_ETRANGERE_SARL / _AU
  - PV_FERMETURE_SUCCURSALE_SARL / _AU

Traitement identique aux Phases 4 / A / B / C (script réutilisable), avec deux garde-fous
étendus propres aux modèles succursale :
  1. Retrait du PRÉAMBULE méta : tout ce qui précède le 1er paragraphe d'en-tête
     (« $DENOMINATION » pour la société marocaine ; « $SOCIETE_MERE_DENOMINATION » pour la
     société mère étrangère) — en-tête « MODÈLE DÉTERMINISTE », Réf/Workflow/Emploi/Base,
     CONVENTION DE BALISAGE, légende, règle d'uniformité, titre méta « EN-TÊTE ».
  2. Retrait de l'ANNEXE glossaire : paragraphe « DICTIONNAIRE DES VARIABLES … » / « ANNEXE »
     et tout ce qui suit (sectPr final conservé).
  2 bis. Retrait du bloc de GUIDAGE de rédaction « Résolutions types … » / « Décisions
     types … » situé entre la boucle RESOLUTIONS et la section CLÔTURE : c'est une consigne
     de rédaction (au même titre que la légende), pas une clause du PV — les résolutions
     réelles sont produites par le mapper via la boucle RESOLUTIONS. Retiré du paragraphe de
     tête jusqu'au titre de section (Heading1) suivant.
  3. Restyle charte JURIKA (Times, en-tête société centré navy, titres de section
     JurikaTitreArticle, corps justifié, marges 2,5 cm, pied numéroté) — le contenu
     juridique directeur n'est JAMAIS modifié. Réexécutable (recopie depuis la source).
"""
import os
import io
import sys
import copy
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HOME = os.path.expanduser("~")
DESKTOP_MODEL = os.path.join(HOME, "Desktop", "model")
DOCX_DIR = os.path.abspath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "backend-java", "ai-service", "src", "main", "resources", "templates", "docx",
))

# Style shell (Times docDefaults + JurikaTitreArticle) transplanté depuis un PV Phase 4.
SHELL_PV = os.path.join(DOCX_DIR, "PV_DEFAUT_QUORUM_SARL.docx")

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x66, 0x66, 0x66)

DENOM_PREFIXES = ("$DENOMINATION", "$SOCIETE_MERE_DENOMINATION")
SOCIETE_PREFIXES = (
    "SOCIÉTÉ À RESPONSABILITÉ LIMITÉE",
    "AU CAPITAL DE",
    "SIÈGE SOCIAL",
    "IMMATRICULÉE AU REGISTRE",
    "RC N°",
    "$SOCIETE_MERE_FORME",
)
DOCTITLE_PREFIXES = (
    "PROCÈS-VERBAL",
    "— CRÉATION",
    "— OUVERTURE",
    "— FERMETURE",
)
CONTROL_PREFIXES = ("◇", "◆", "▼", "▲", "▶", "◁")
STRIP_ANNEX_PREFIXES = ("DICTIONNAIRE DES VARIABLES", "ANNEXE")
GUIDANCE_PREFIXES = ("résolutions types", "resolutions types",
                     "décisions types", "decisions types")


def para_text(p):
    el = getattr(p, "_p", p)
    return "".join(n.text or "" for n in el.iter(qn("w:t")))


def source_style_id(p):
    ppr = p._p.find(qn("w:pPr"))
    if ppr is None:
        return None
    ps = ppr.find(qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


def _child_style_id(ch):
    ppr = ch.find(qn("w:pPr"))
    if ppr is None:
        return None
    ps = ppr.find(qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


# ----------------------------------------------------------------------------
# 1 + 2 : strip préambule / annexe glossaire (niveau body, préserve sectPr)
# ----------------------------------------------------------------------------
def strip_preamble_and_annex(doc, start_marker):
    body = doc.element.body
    children = list(body.iterchildren())

    start = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p") and para_text(ch).strip().startswith(start_marker):
            start = i
            break
    if start is None:
        raise RuntimeError("Marqueur %s introuvable — préambule non retiré." % start_marker)

    annex = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p"):
            t = para_text(ch).strip()
            if any(t.startswith(pref) for pref in STRIP_ANNEX_PREFIXES):
                annex = i
                break

    sectpr = None
    if children and children[-1].tag == qn("w:sectPr"):
        sectpr = children[-1]

    end = annex if annex is not None else len(children)
    keep = children[start:end]

    for ch in children:
        body.remove(ch)
    for ch in keep:
        body.append(ch)
    if sectpr is not None:
        body.append(sectpr)


# ----------------------------------------------------------------------------
# 2 bis : strip bloc guidage « Résolutions/Décisions types … » (jusqu'au Heading1 suivant)
# ----------------------------------------------------------------------------
def strip_guidance_block(doc):
    body = doc.element.body
    children = list(body.iterchildren())

    start = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p"):
            t = para_text(ch).strip().lower()
            if any(t.startswith(pref) for pref in GUIDANCE_PREFIXES):
                start = i
                break
    if start is None:
        return False  # tolérant : modèle sans bloc de guidage

    end = len(children)
    for i in range(start + 1, len(children)):
        ch = children[i]
        if ch.tag == qn("w:p") and _child_style_id(ch) and _child_style_id(ch).startswith("Heading"):
            end = i
            break

    for ch in children[start:end]:
        body.remove(ch)
    return True


# ----------------------------------------------------------------------------
# 3 : styles.xml — docDefaults Times + style JurikaTitreArticle (transplanté)
# ----------------------------------------------------------------------------
def ensure_doc_defaults_times(doc):
    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        dd = OxmlElement("w:docDefaults")
        styles_el.insert(0, dd)
    rprd = dd.find(qn("w:rPrDefault"))
    if rprd is None:
        rprd = OxmlElement("w:rPrDefault")
        dd.insert(0, rprd)
    rpr = rprd.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rprd.append(rpr)
    el = rpr.find(qn("w:rFonts"))
    if el is not None:
        rpr.remove(el)
    rfonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    rpr.insert(0, rfonts)
    for tag, val in (("w:sz", "22"), ("w:szCs", "22")):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), val)


def ensure_jurika_titre_style(doc):
    styles_el = doc.styles.element
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "JurikaTitreArticle":
            return
    shell = Document(SHELL_PV)
    src = None
    for st in shell.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "JurikaTitreArticle":
            src = st
            break
    if src is None:
        raise RuntimeError("JurikaTitreArticle absent du shell PV.")
    styles_el.append(copy.deepcopy(src))


# ----------------------------------------------------------------------------
# 3 : restyle runs + paragraphes
# ----------------------------------------------------------------------------
def set_run_font(run, color):
    run.font.name = "Times New Roman"
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    if color is not None:
        run.font.color.rgb = color


def set_pstyle(p, style_id):
    ppr = p._p.get_or_add_pPr()
    ps = ppr.find(qn("w:pStyle"))
    if ps is None:
        ps = OxmlElement("w:pStyle")
        ppr.insert(0, ps)
    ps.set(qn("w:val"), style_id)


def classify(p):
    if source_style_id(p) == "Heading1":
        return "section"
    t = para_text(p).strip()
    if not t:
        return "body"
    if t.startswith(CONTROL_PREFIXES):
        return "control"
    if any(t.startswith(pref) for pref in DENOM_PREFIXES):
        return "denom"
    if any(t.startswith(pref) for pref in DOCTITLE_PREFIXES):
        return "doctitle"
    if any(t.startswith(pref) for pref in SOCIETE_PREFIXES):
        return "societe"
    return "body"


def restyle_paragraph(p):
    kind = classify(p)
    if kind == "denom":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            set_run_font(r, NAVY)
    elif kind == "societe":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = False
            set_run_font(r, BLACK)
    elif kind == "doctitle":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            set_run_font(r, BLACK)
    elif kind == "section":
        set_pstyle(p, "JurikaTitreArticle")
        for r in p.runs:
            set_run_font(r, NAVY)
    elif kind == "control":
        for r in p.runs:
            set_run_font(r, BLACK)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            set_run_font(r, BLACK)


def restyle_all(doc):
    for p in doc.paragraphs:
        restyle_paragraph(p)


# ----------------------------------------------------------------------------
# 3 : marges + pied de page numéroté
# ----------------------------------------------------------------------------
def set_margins_and_footer(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)
        footer = s.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        set_run_font(run, GREY)
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(instr); run._r.append(f2)


# ----------------------------------------------------------------------------
# Orchestration
# ----------------------------------------------------------------------------
def integrate(src_name, dest_code, start_marker, must_contain):
    src_path = os.path.join(DESKTOP_MODEL, src_name)
    dest = os.path.join(DOCX_DIR, dest_code + ".docx")
    if not os.path.exists(src_path):
        raise RuntimeError("Source introuvable : " + src_path)
    shutil.copyfile(src_path, dest)
    doc = Document(dest)
    strip_preamble_and_annex(doc, start_marker)
    guidance_removed = strip_guidance_block(doc)
    ensure_doc_defaults_times(doc)
    ensure_jurika_titre_style(doc)
    restyle_all(doc)
    set_margins_and_footer(doc)
    doc.save(dest)

    check = Document(dest)
    full = "\n".join(para_text(p) for p in check.paragraphs)
    assert "CONVENTION DE BALISAGE" not in full, "préambule résiduel dans " + dest_code
    assert "DICTIONNAIRE DES VARIABLES" not in full, "annexe glossaire résiduelle dans " + dest_code
    assert "MODÈLE DÉTERMINISTE" not in full, "en-tête méta résiduel dans " + dest_code
    assert "types (à renseigner" not in full, "bloc de guidage résiduel dans " + dest_code
    assert start_marker in full, "marqueur %s perdu dans %s" % (start_marker, dest_code)
    for v in must_contain:
        assert v in full, "variable %s perdue dans %s" % (v, dest_code)
    print("  OK -> %s (%d paragraphes, guidage retiré=%s)"
          % (dest_code, len(check.paragraphs), guidance_removed))


def main():
    print("Intégration modèles Succursales dans:", DOCX_DIR)
    integrate("PV_CREATION_SUCCURSALE_MAROC_SARL.docx",
              "PV_CREATION_SUCCURSALE_MAROC_SARL", "$DENOMINATION",
              ("$SUCCURSALE_ENSEIGNE", "$SUCCURSALE_DOTATION_PRESENTE",
               "$SUCCURSALE_RESPONSABLE_PRESENT", "$SUCCURSALE_VILLE_GREFFE"))
    integrate("PV_CREATION_SUCCURSALE_MAROC_SARL_AU.docx",
              "PV_CREATION_SUCCURSALE_MAROC_SARL_AU", "$DENOMINATION",
              ("$SUCCURSALE_ENSEIGNE", "$SUCCURSALE_DOTATION_PRESENTE",
               "$GERANT_NOM", "$ASSOCIE_TYPE"))
    integrate("PV_CREATION_SUCCURSALE_ETRANGERE_SARL.docx",
              "PV_CREATION_SUCCURSALE_ETRANGERE_SARL", "$SOCIETE_MERE_DENOMINATION",
              ("$SOCIETE_MERE_DENOMINATION", "$ORGANE_COMPETENT",
               "$ORGANE_DECISION_DATE", "$SUCCURSALE_ENSEIGNE"))
    integrate("PV_CREATION_SUCCURSALE_ETRANGERE_SARL_AU.docx",
              "PV_CREATION_SUCCURSALE_ETRANGERE_SARL_AU", "$SOCIETE_MERE_DENOMINATION",
              ("$SOCIETE_MERE_DENOMINATION", "$ORGANE_COMPETENT", "$SUCCURSALE_ENSEIGNE"))
    integrate("PV_FERMETURE_SUCCURSALE_SARL.docx",
              "PV_FERMETURE_SUCCURSALE_SARL", "$DENOMINATION",
              ("$SUCCURSALE_RC_NUMERO", "$SUCCURSALE_DATE_FERMETURE",
               "$SUCCURSALE_MOTIF", "$SUCCURSALE_RESPONSABLE_PRESENT"))
    integrate("PV_FERMETURE_SUCCURSALE_SARL_AU.docx",
              "PV_FERMETURE_SUCCURSALE_SARL_AU", "$DENOMINATION",
              ("$SUCCURSALE_RC_NUMERO", "$SUCCURSALE_DATE_FERMETURE",
               "$SUCCURSALE_MOTIF", "$ASSOCIE_TYPE"))
    print("Terminé.")


if __name__ == "__main__":
    main()
