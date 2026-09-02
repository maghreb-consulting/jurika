# -*- coding: utf-8 -*-
"""
Lot « Liquidation 4 étapes » (2026-08-13) — intégration des 2 modèles directeur
« annonce légale de CLÔTURE DE LIQUIDATION » :

  - ANNONCE_LEGALE_LIQUIDATION_SARL     (source : 04_LIQUIDATION/ANNONCE_LEGALE_LIQUIDATION_SARL_modele_deterministe.docx)
  - ANNONCE_LEGALE_LIQUIDATION_SARL_AU  (source : 04_LIQUIDATION/ANNONCE_LEGALE_LIQUIDATION_SARL_AU_modele_deterministe.docx)

Même traitement que l'annonce de DISSOLUTION (2026-08-12) :

  1. On ne conserve QUE l'avis publiable : du titre « AVIS DE CLÔTURE DE LIQUIDATION »
     jusqu'au paragraphe « Le dépôt légal … » inclus. Tout le scaffolding NON publié
     (en-tête « MODÈLE DÉTERMINISTE », Réf./Workflow/Objet/Base légale, « CONVENTION DE
     BALISAGE », « NOTES D'EMPLOI ») est retiré.
  2. Le CONTENU juridique du directeur est repris **verbatim** (texte des runs + gras
     d'origine + marqueurs $VAR et ◇/◆) — jamais reformulé.
  3. La FORME est celle de la charte JURIKA déjà appliquée à l'annonce de dissolution :
     coquille = ANNONCE_LEGALE_DISSOLUTION_SARL.docx (styles.xml, marges, sectPr), titre
     Times 16 pt gras navy 1F4E79, corps Times 11 pt noir, spacing after 60.

Le script est **réexécutable** : il repart à chaque fois de la source directeur.
Il régénère aussi les baselines JSON consommées par DirecteurTemplateContentUnchangedTest.

Usage : python scripts/integrate-liquidation-annonce-templates-2026-08-13.py
"""
import copy
import io
import json
import os
import shutil
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HOME = os.path.expanduser("~")
SRC_DIR = os.path.join(HOME, "Desktop", "model", "livrable deterministe", "04_LIQUIDATION")
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DOCX_DIR = os.path.join(ROOT, "backend-java", "ai-service", "src", "main",
                        "resources", "templates", "docx")
BASELINE_DIR = os.path.join(ROOT, "backend-java", "ai-service", "src", "test",
                            "resources", "directeur-baseline")

# Coquille de forme : l'annonce de dissolution, déjà à la charte (Times, marges, sectPr).
SHELL = os.path.join(DOCX_DIR, "ANNONCE_LEGALE_DISSOLUTION_SARL.docx")

TITRE = "AVIS DE CLÔTURE DE LIQUIDATION"
FIN_PREFIX = "Le dépôt légal"

TITRE_SZ = "32"   # demi-points -> 16 pt
CORPS_SZ = "22"   # demi-points -> 11 pt
NAVY = "1F4E79"
NOIR = "000000"


def para_runs(p_el):
    """[(gras, texte)] des runs non vides d'un <w:p>."""
    out = []
    for r in p_el.iter(qn("w:r")):
        text = "".join(t.text or "" for t in r.iter(qn("w:t")))
        if not text:
            continue
        rpr = r.find(qn("w:rPr"))
        bold = False
        if rpr is not None:
            b = rpr.find(qn("w:b"))
            bold = b is not None and b.get(qn("w:val")) not in ("0", "false")
        out.append((bold, text))
    return out


def para_text(p_el):
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def extract_avis(src_path):
    """Paragraphes de l'avis publiable : [[(gras, texte), ...], ...] (blancs inclus)."""
    doc = Document(src_path)
    children = [c for c in doc.element.body.iterchildren() if c.tag == qn("w:p")]

    start = next((i for i, c in enumerate(children) if para_text(c).strip() == TITRE), None)
    if start is None:
        raise RuntimeError("Titre « %s » introuvable dans %s" % (TITRE, src_path))
    end = next((i for i, c in enumerate(children)
                if i > start and para_text(c).strip().startswith(FIN_PREFIX)), None)
    if end is None:
        raise RuntimeError("Paragraphe « %s … » introuvable dans %s" % (FIN_PREFIX, src_path))

    return [para_runs(c) for c in children[start:end + 1]]


def make_run(text, bold, size, color):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rpr.append(rfonts)
    b = OxmlElement("w:b")
    if not bold:
        b.set(qn("w:val"), "0")
    rpr.append(b)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rpr.append(col)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), size)
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def make_paragraph(runs, is_title):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "60")
    ppr.append(spacing)
    p.append(ppr)
    for bold, text in runs:
        if is_title:
            p.append(make_run(text, True, TITRE_SZ, NAVY))
        else:
            p.append(make_run(text, bold, CORPS_SZ, NOIR))
    return p


def build(src_name, dest_code):
    src_path = os.path.join(SRC_DIR, src_name)
    if not os.path.exists(src_path):
        raise RuntimeError("Source directeur introuvable : " + src_path)
    if not os.path.exists(SHELL):
        raise RuntimeError("Coquille de forme introuvable : " + SHELL)

    paragraphs = extract_avis(src_path)

    dest = os.path.join(DOCX_DIR, dest_code + ".docx")
    shutil.copyfile(SHELL, dest)
    doc = Document(dest)
    body = doc.element.body

    sectpr = body.find(qn("w:sectPr"))
    sectpr = copy.deepcopy(sectpr) if sectpr is not None else None
    for ch in list(body.iterchildren()):
        body.remove(ch)
    for idx, runs in enumerate(paragraphs):
        body.append(make_paragraph(runs, is_title=(idx == 0)))
    if sectpr is not None:
        body.append(sectpr)
    doc.save(dest)

    # ---- Garde-fous : contenu directeur intact, scaffolding absent ----
    check = Document(dest)
    lines = [p.text for p in check.paragraphs]
    full = "\n".join(lines)
    for banned in ("MODÈLE DÉTERMINISTE", "CONVENTION DE BALISAGE",
                   "NOTES D'EMPLOI", "Réf. modèle", "Base légale"):
        assert banned not in full, "scaffolding résiduel (%s) dans %s" % (banned, dest_code)
    for var in ("$DENOMINATION", "$CAPITAL_CHIFFRES", "$SIEGE_SOCIAL", "$RC_NUMERO",
                "$VILLE_GREFFE", "$DATE_CLOTURE_LIQUIDATION", "$RESULTAT_LIQUIDATION_TYPE",
                "$BONI_LIQUIDATION_CHIFFRES", "$MALI_LIQUIDATION_CHIFFRES",
                "$LIQUIDATEUR_CIVILITE", "$LIQUIDATEUR_PRENOM", "$LIQUIDATEUR_NOM",
                "$DATE_DEPOT_LEGAL", "$DEPOT_LEGAL_NUMERO"):
        assert var in full, "variable %s perdue dans %s" % (var, dest_code)
    for marker in ("◇ SI :", "◇ SINON", "◆ FIN SI"):
        assert marker in full, "marqueur %s perdu dans %s" % (marker, dest_code)
    assert lines[0].strip() == TITRE, "titre inattendu dans %s : %r" % (dest_code, lines[0])

    # ---- Baseline JSON (garde-fou « contenu intouchable ») ----
    os.makedirs(BASELINE_DIR, exist_ok=True)
    baseline_path = os.path.join(BASELINE_DIR, dest_code + ".docx.json")
    with open(baseline_path, "w", encoding="utf-8") as fh:
        json.dump(lines, fh, ensure_ascii=False, indent=2)

    print("  OK -> %s (%d paragraphes) + baseline" % (dest_code, len(lines)))


def main():
    print("Intégration annonces de clôture de liquidation dans :", DOCX_DIR)
    build("ANNONCE_LEGALE_LIQUIDATION_SARL_modele_deterministe.docx",
          "ANNONCE_LEGALE_LIQUIDATION_SARL")
    build("ANNONCE_LEGALE_LIQUIDATION_SARL_AU_modele_deterministe.docx",
          "ANNONCE_LEGALE_LIQUIDATION_SARL_AU")
    print("Terminé.")


if __name__ == "__main__":
    main()
