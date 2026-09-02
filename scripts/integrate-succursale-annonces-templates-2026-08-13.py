# -*- coding: utf-8 -*-
"""
Lot DIVERS (2026-08-13) — intégration des 4 modèles directeur « annonce légale
de SUCCURSALE » :

  - ANNONCE_LEGALE_OUVERTURE_SUCCURSALE_SARL     (05_SUCCURSALE_OUVERTURE)
  - ANNONCE_LEGALE_OUVERTURE_SUCCURSALE_SARL_AU  (05_SUCCURSALE_OUVERTURE)
  - ANNONCE_LEGALE_FERMETURE_SUCCURSALE_SARL     (06_SUCCURSALE_FERMETURE)
  - ANNONCE_LEGALE_FERMETURE_SUCCURSALE_SARL_AU  (06_SUCCURSALE_FERMETURE)

Traitement identique aux annonces de DISSOLUTION (2026-08-12) et de CLÔTURE DE
LIQUIDATION (2026-08-13) :

  1. On ne conserve QUE l'avis publiable : du titre « AVIS D'OUVERTURE DE
     SUCCURSALE » / « AVIS DE FERMETURE DE SUCCURSALE » jusqu'au paragraphe
     « Le dépôt légal … » inclus. Tout le scaffolding NON publié (en-tête
     « MODÈLE DÉTERMINISTE », Réf./Workflow/Objet/Base légale, « CONVENTION DE
     BALISAGE », « NOTES D'EMPLOI ») est retiré.
  2. Le CONTENU juridique du directeur est repris **verbatim** (texte des runs +
     gras d'origine + marqueurs $VAR et ◇/◆) — jamais reformulé.
  3. La FORME est celle de la charte JURIKA déjà appliquée aux autres annonces :
     coquille = ANNONCE_LEGALE_DISSOLUTION_SARL.docx (styles.xml, marges, sectPr),
     titre Times 16 pt gras navy 1F4E79, corps Times 11 pt noir, spacing after 60.

Le script est **réexécutable** : il repart à chaque fois de la source directeur.
Il régénère aussi les baselines JSON consommées par
DirecteurTemplateContentUnchangedTest.

Usage : python scripts/integrate-succursale-annonces-templates-2026-08-13.py
"""
import copy
import io
import json
import os
import shutil
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HOME = os.path.expanduser("~")
MODELS = os.path.join(HOME, "Desktop", "model", "livrable deterministe")
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DOCX_DIR = os.path.join(ROOT, "backend-java", "ai-service", "src", "main",
                        "resources", "templates", "docx")
BASELINE_DIR = os.path.join(ROOT, "backend-java", "ai-service", "src", "test",
                            "resources", "directeur-baseline")

# Coquille de forme : l'annonce de dissolution, déjà à la charte (Times, marges, sectPr).
SHELL = os.path.join(DOCX_DIR, "ANNONCE_LEGALE_DISSOLUTION_SARL.docx")

FIN_PREFIX = "Le dépôt légal"

TITRE_SZ = "32"   # demi-points -> 16 pt
CORPS_SZ = "22"   # demi-points -> 11 pt
NAVY = "1F4E79"
NOIR = "000000"

# (dossier source, fichier source, code destination, mot-clé attendu dans le titre)
JOBS = [
    ("05_SUCCURSALE_OUVERTURE",
     "ANNONCE_LEGALE_OUVERTURE_SUCCURSALE_SARL_modele_deterministe.docx",
     "ANNONCE_LEGALE_OUVERTURE_SUCCURSALE_SARL", "OUVERTURE"),
    ("05_SUCCURSALE_OUVERTURE",
     "ANNONCE_LEGALE_OUVERTURE_SUCCURSALE_SARL_AU_modele_deterministe.docx",
     "ANNONCE_LEGALE_OUVERTURE_SUCCURSALE_SARL_AU", "OUVERTURE"),
    ("06_SUCCURSALE_FERMETURE",
     "ANNONCE_LEGALE_FERMETURE_SUCCURSALE_SARL_modele_deterministe.docx",
     "ANNONCE_LEGALE_FERMETURE_SUCCURSALE_SARL", "FERMETURE"),
    ("06_SUCCURSALE_FERMETURE",
     "ANNONCE_LEGALE_FERMETURE_SUCCURSALE_SARL_AU_modele_deterministe.docx",
     "ANNONCE_LEGALE_FERMETURE_SUCCURSALE_SARL_AU", "FERMETURE"),
]

# Variables dont la présence est vérifiée après reconstruction, par famille d'avis.
REQUIRED_VARS = {
    "OUVERTURE": ("$DENOMINATION", "$CAPITAL_CHIFFRES", "$SIEGE_SOCIAL", "$RC_NUMERO",
                  "$VILLE_GREFFE", "$ASSEMBLEE_DATE", "$SUCCURSALE_ENSEIGNE",
                  "$SUCCURSALE_ADRESSE", "$SUCCURSALE_VILLE", "$SUCCURSALE_ACTIVITE",
                  "$SUCCURSALE_DATE_OUVERTURE", "$SUCCURSALE_VILLE_GREFFE",
                  "$SUCCURSALE_DOTATION_PRESENTE", "$SUCCURSALE_DOTATION_CHIFFRES",
                  "$SUCCURSALE_DOTATION_LETTRES", "$SUCCURSALE_RESPONSABLE_PRESENT",
                  "$SUCCURSALE_RESPONSABLE_CIVILITE", "$SUCCURSALE_RESPONSABLE_PRENOM",
                  "$SUCCURSALE_RESPONSABLE_NOM", "$SUCCURSALE_RESPONSABLE_POUVOIRS",
                  "$DATE_DEPOT_LEGAL", "$DEPOT_LEGAL_NUMERO"),
    "FERMETURE": ("$DENOMINATION", "$CAPITAL_CHIFFRES", "$SIEGE_SOCIAL", "$RC_NUMERO",
                  "$VILLE_GREFFE", "$ASSEMBLEE_DATE", "$SUCCURSALE_ENSEIGNE",
                  "$SUCCURSALE_ADRESSE", "$SUCCURSALE_VILLE", "$SUCCURSALE_VILLE_GREFFE",
                  "$SUCCURSALE_RC_NUMERO", "$SUCCURSALE_DATE_FERMETURE",
                  "$SUCCURSALE_MOTIF", "$DATE_DEPOT_LEGAL", "$DEPOT_LEGAL_NUMERO"),
}

# L'avis de fermeture est linéaire (aucun bloc conditionnel) ; seul l'avis
# d'ouverture porte les deux blocs ◇ SI (dotation / responsable).
REQUIRED_MARKERS = {
    "OUVERTURE": ("◇ SI :", "◆ FIN SI"),
    "FERMETURE": (),
}


def para_runs(p_el):
    """[(gras, texte)] des runs non vides d'un <w:p>."""
    out = []
    for r in p_el.iter(qn("w:r")):
        text = "".join(t.text or "" for t in r.iter(qn("w:t")))
        if not text:
            continue
        rpr = r.find(qn("w:rPr"))
        bold = False
        if rpr is not None:
            b = rpr.find(qn("w:b"))
            bold = b is not None and b.get(qn("w:val")) not in ("0", "false")
        out.append((bold, text))
    return out


def para_text(p_el):
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def find_title(children, keyword):
    """
    Index du titre de l'avis publiable. On ne compare PAS à une chaîne littérale :
    l'apostrophe du directeur peut être droite (U+0027) ou typographique (U+2019).
    On cherche le paragraphe qui commence par « AVIS D » et contient le mot-clé
    (OUVERTURE / FERMETURE) ainsi que « SUCCURSALE ».
    """
    for i, c in enumerate(children):
        t = para_text(c).strip()
        if t.startswith("AVIS D") and keyword in t and "SUCCURSALE" in t:
            return i
    return None


def extract_avis(src_path, keyword):
    """Paragraphes de l'avis publiable : [[(gras, texte), ...], ...] (blancs inclus)."""
    doc = Document(src_path)
    children = [c for c in doc.element.body.iterchildren() if c.tag == qn("w:p")]

    start = find_title(children, keyword)
    if start is None:
        raise RuntimeError("Titre « AVIS D…%s…SUCCURSALE » introuvable dans %s"
                           % (keyword, src_path))
    end = next((i for i, c in enumerate(children)
                if i > start and para_text(c).strip().startswith(FIN_PREFIX)), None)
    if end is None:
        raise RuntimeError("Paragraphe « %s … » introuvable dans %s" % (FIN_PREFIX, src_path))

    return [para_runs(c) for c in children[start:end + 1]]


def make_run(text, bold, size, color):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rpr.append(rfonts)
    b = OxmlElement("w:b")
    if not bold:
        b.set(qn("w:val"), "0")
    rpr.append(b)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rpr.append(col)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), size)
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def make_paragraph(runs, is_title):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "60")
    ppr.append(spacing)
    p.append(ppr)
    for bold, text in runs:
        if is_title:
            p.append(make_run(text, True, TITRE_SZ, NAVY))
        else:
            p.append(make_run(text, bold, CORPS_SZ, NOIR))
    return p


def build(src_dir, src_name, dest_code, keyword):
    src_path = os.path.join(MODELS, src_dir, src_name)
    if not os.path.exists(src_path):
        raise RuntimeError("Source directeur introuvable : " + src_path)
    if not os.path.exists(SHELL):
        raise RuntimeError("Coquille de forme introuvable : " + SHELL)

    paragraphs = extract_avis(src_path, keyword)

    dest = os.path.join(DOCX_DIR, dest_code + ".docx")
    shutil.copyfile(SHELL, dest)
    doc = Document(dest)
    body = doc.element.body

    sectpr = body.find(qn("w:sectPr"))
    sectpr = copy.deepcopy(sectpr) if sectpr is not None else None
    for ch in list(body.iterchildren()):
        body.remove(ch)
    for idx, runs in enumerate(paragraphs):
        body.append(make_paragraph(runs, is_title=(idx == 0)))
    if sectpr is not None:
        body.append(sectpr)
    doc.save(dest)

    # ---- Garde-fous : contenu directeur intact, scaffolding absent ----
    check = Document(dest)
    lines = [p.text for p in check.paragraphs]
    full = "\n".join(lines)
    for banned in ("MODÈLE DÉTERMINISTE", "CONVENTION DE BALISAGE",
                   "NOTES D'EMPLOI", "NOTES D’EMPLOI", "Réf. modèle", "Base légale",
                   "$NOM_EN_MAJUSCULES"):
        assert banned not in full, "scaffolding résiduel (%s) dans %s" % (banned, dest_code)
    for var in REQUIRED_VARS[keyword]:
        assert var in full, "variable %s perdue dans %s" % (var, dest_code)
    for marker in REQUIRED_MARKERS[keyword]:
        assert marker in full, "marqueur %s perdu dans %s" % (marker, dest_code)
    assert lines[0].strip().startswith("AVIS D"), \
        "titre inattendu dans %s : %r" % (dest_code, lines[0])

    # ---- Baseline JSON (garde-fou « contenu intouchable ») ----
    os.makedirs(BASELINE_DIR, exist_ok=True)
    baseline_path = os.path.join(BASELINE_DIR, dest_code + ".docx.json")
    with open(baseline_path, "w", encoding="utf-8") as fh:
        json.dump(lines, fh, ensure_ascii=False, indent=2)

    print("  OK -> %s (%d paragraphes) + baseline" % (dest_code, len(lines)))


def main():
    print("Intégration annonces succursale (ouverture + fermeture) dans :", DOCX_DIR)
    for src_dir, src_name, dest_code, keyword in JOBS:
        build(src_dir, src_name, dest_code, keyword)
    print("Terminé.")


if __name__ == "__main__":
    main()
