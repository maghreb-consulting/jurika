#!/usr/bin/env python3
"""
Sprint 2026-06-12 — Conversion des annotations NL des templates Création SARL
en marqueurs machine-lisibles ◇ FLAG / ◆ FLAG pour le moteur DocxTemplateEngine.

Cible : 6 templates dans backend-java/ai-service/src/main/resources/templates/docx/
  - STATUTS_CONSTITUTIFS_SARL.docx
  - STATUTS_CONSTITUTIFS_SARL_AU.docx
  - ANNONCE_JAL_SARL.docx
  - ANNONCE_JAL_SARL_AU.docx
  - ACTE_NOMINATION_GERANT_SARL.docx
  - ACTE_NOMINATION_GERANT_SARL_AU.docx

Transformations :
  - "(Si {{flag}} = oui) ..." → ◇ FLAG (préfixe paragraphe) + ◆ FLAG (suffixe)
    et suppression du préfixe NL "(Si {{flag}} = oui)".
  - "(Si applicable) ..." → wrap conditionnel sur flag déduit du contexte
    (apport_nature / fonds_commerce_apport / apport_industrie).
  - Variantes inline "(Variante 1, ...)/(Variante 2, ...)" → remplacement par
    placeholder unique {{...clause}} dont le mapper choisit la valeur.
  - "(Option tacite reconduction, à activer si {{duree_modalite}} = ...)"
    → remplacement par placeholder {{duree_clause_tacite}} (le mapper choisit).
  - "(Boucle) / (etc.) / annotations légères" → retirées (le moteur ▶ gère
    déjà la répétition au niveau de l'article 7).

Mode opératoire : on lit le texte complet de chaque paragraphe via la
concaténation des runs (insensible aux fragmentations Word), on substitue
en mémoire, puis on réécrit dans le 1er run et on vide les autres
(comportement aligné avec le moteur Java — limitation L1 documentée).
Sauvegarde la copie originale à côté avec suffixe `.preconv.docx` pour
audit (gardée hors-git via .gitignore).
"""

from __future__ import annotations

import io
import os
import re
import shutil
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

try:
    from docx import Document
except ImportError:
    print("python-docx requis : pip install python-docx", file=sys.stderr)
    sys.exit(1)


ROOT = Path(__file__).resolve().parent.parent
TPL_DIR = ROOT / "backend-java" / "ai-service" / "src" / "main" / "resources" / "templates" / "docx"


# =============================================================================
# Helpers paragraph rewrite
# =============================================================================

def para_text(para) -> str:
    return "".join(run.text or "" for run in para.runs)


def rewrite_para_text(para, new_text: str) -> None:
    """Réécrit le texte d'un paragraphe en gardant le style du 1er run.
    Vide les runs suivants. Limitation L1 acceptée (alignée moteur Java)."""
    runs = para.runs
    if not runs:
        # Paragraphe sans run — ajouter un nouveau run préserve les properties.
        para.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


# =============================================================================
# Mutations communes
# =============================================================================

# Pattern : « (Si {{flag_name}} = oui) » — capture le flag.
RE_COND_OUI = re.compile(r"\(\s*Si\s+\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}\s*=\s*oui\s*\)\s*")

# Pattern : « (Si {{flag}} = oui et ...) » avec parenthèse étendue.
RE_COND_OUI_EXT = re.compile(r"\(\s*Si\s+\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}\s*=\s*oui[^)]*\)\s*")

# Pattern : « (Si applicable...) » (sans flag explicite, le flag est déduit).
RE_SI_APPLICABLE = re.compile(r"\(\s*Si\s+applicable[^)]*\)\s*", re.IGNORECASE)


def wrap_cond_flag(text: str, flag: str) -> str:
    """Préfixe ◇ FLAG, suffixe ◆ FLAG sur un texte (paragraphe entier)."""
    flag_up = flag.upper()
    return f"◇ {flag_up} {text.strip()} ◆ {flag_up}"


# =============================================================================
# Spec de transformation par template
# =============================================================================
#
# Chaque entrée = match (substring distinguant le paragraphe) + transform.
# Si transform est str : remplace le paragraphe entier par cette string.
# Si transform est callable : reçoit le texte et retourne le nouveau texte.

def transform_si_oui(text: str) -> str:
    """Détecte le flag dans `(Si {{flag}} = oui)` et wrap le paragraphe entier."""
    m = RE_COND_OUI_EXT.search(text)
    if not m:
        m = RE_COND_OUI.search(text)
    if not m:
        return text
    flag = m.group(1)
    # On nettoie aussi les éventuelles parenthèses inline qui suivent.
    cleaned = RE_COND_OUI_EXT.sub("", text)
    cleaned = RE_COND_OUI.sub("", cleaned)
    return wrap_cond_flag(cleaned, flag)


def wrap_flag(flag: str):
    def inner(text: str) -> str:
        cleaned = RE_SI_APPLICABLE.sub("", text)
        cleaned = RE_COND_OUI_EXT.sub("", cleaned)
        cleaned = RE_COND_OUI.sub("", cleaned)
        return wrap_cond_flag(cleaned, flag)
    return inner


def clean_inline_annotations(text: str) -> str:
    """Retire les annotations parasites résiduelles qui ne sont pas des conditions :
    « (Boucle) », « (etc.) », « (Option ...) ». Conserve placeholder du mapper."""
    s = text
    # « (Boucle) … {{associe_pp_nom_prenom}} : … parts ; (etc.) … » — ART 7
    s = re.sub(r"\(\s*Boucle\s*\)\s*", "", s)
    s = re.sub(r"\s*;\s*\(\s*etc\.?\s*\)\s*", " ", s)
    # « (Option tacite reconduction, à activer si {{duree_modalite}} = ...) »
    # On le remplace par un placeholder dédié — le mapper injectera le texte ad-hoc.
    s = re.sub(
        r"\(\s*Option\s+tacite[^)]*?{{\s*duree_modalite\s*}}[^)]*\)",
        "{{duree_clause_tacite}}",
        s,
        flags=re.DOTALL,
    )
    # === Collapse des blocs « (Variante X) ... (Variante Y) ... » mutuellement
    # exclusifs en un placeholder unique résolu par le mapper. ===
    #
    # SARL art 6.2 — apports en nature : variantes commissaire / associés.
    s = re.sub(
        r"L'évaluation est faite\s*\{\{\s*apports_nature_evaluation\s*\}\}\s*:[\s\S]*?valeur attribuée aux apports en nature\.",
        "L'évaluation est faite selon les modalités suivantes : {{apports_nature_evaluation_clause}}",
        s,
    )
    # SARL_AU art 6.2 — apports en nature : variantes commissaire / associé unique.
    s = re.sub(
        r"L'évaluation est faite\s*\{\{\s*apport_nature_evaluation\s*\}\}\s*:[\s\S]*?valeur attribuée aux apports en nature\.",
        "L'évaluation est faite selon les modalités suivantes : {{apport_nature_evaluation_clause}}",
        s,
    )
    # SARL — engagement société : variantes signature gérant.
    s = re.sub(
        r"La société est engagée vis-à-vis des tiers par\s*\{\{\s*gerant_signature_mode\s*\}\}\s*\(Variante[\s\S]*?personne morale\)\.",
        "La société est engagée vis-à-vis des tiers selon les modalités suivantes : {{gerant_signature_mode_clause}}.",
        s,
    )
    # SARL — juridiction : variantes tribunal compétent / arbitrage (1 seule paire
    # de parenthèses contenant Variante 1 ; Variante 2).
    s = re.sub(
        r"Toutes contestations[\s\S]*?\(Variante 1\s*:[\s\S]*?Variante 2\s*:[\s\S]*?sans appel\)\.",
        "Toutes contestations relatives aux affaires sociales sont réglées selon les modalités suivantes : {{clause_juridiction_clause}}.",
        s,
    )
    # SARL_AU art 14.2 — désignation initiale de la gérance : variantes A/B/C/D.
    s = re.sub(
        r"14\.2 Désignation initiale[\s\S]*?représentant permanent\s*\{\{\s*gerant_pm_representant_permanent\s*\}\}\.",
        "14.2 Désignation initiale — {{gerance_au_designation_clause}}",
        s,
    )
    # « (si distincts de l'associé unique) » — SARL_AU, signature.
    s = re.sub(
        r"\(\s*si\s+distincts\s+de\s+l['’]associé\s+unique\s*\)\s*",
        "",
        s,
        flags=re.IGNORECASE,
    )
    return s


# Spec — clé = nom de fichier, valeur = liste de (matcher, transform).
# Matcher = string → must be `in` paragraph text.
# Transform = callable(text) -> text, OU str remplaçant entièrement.
SPEC = {
    "STATUTS_CONSTITUTIFS_SARL.docx": [
        # ART 5 Durée — option tacite reconduction inline.
        ("Option tacite reconduction", clean_inline_annotations),
        # ART 6.2 Apports en nature.
        ("6.2 Apports en nature", wrap_flag("APPORT_NATURE")),
        # ART 6.3 Apport de fonds de commerce.
        ("6.3 Apport de fonds de commerce", wrap_flag("FONDS_COMMERCE_APPORT")),
        # ART 6.4 Apport en industrie.
        ("6.4 Apport en industrie", wrap_flag("APPORT_INDUSTRIE")),
        # ART 7 — (Boucle) / (etc.) / le moteur ▶ APPORTS_NUMERAIRE devrait gérer ; on nettoie.
        ("(Boucle)", clean_inline_annotations),
        # ART 12 — Clauses optionnelles.
        ("12.1 Droit de préemption", transform_si_oui),
        ("12.2 Inaliénabilité", transform_si_oui),
        ("12.3 Clause d'exclusion", transform_si_oui),
        ("12.3 Clause d’exclusion", transform_si_oui),
        ("12.4 Sortie conjointe", transform_si_oui),
        ("12.5 Sortie forcée", transform_si_oui),
        ("12.6 Pacte d'associés", transform_si_oui),
        ("12.6 Pacte d’associés", transform_si_oui),
        # Variantes inline → collapsées par clean_inline_annotations.
        ("apports_nature_evaluation", clean_inline_annotations),
        ("gerant_signature_mode", clean_inline_annotations),
        ("Toutes contestations", clean_inline_annotations),
        # CAC (article 17 ou similaire).
        ("Est nommé(e) commissaire aux comptes", transform_si_oui),
        # Acomptes sur dividendes.
        ("acomptes sur dividendes", transform_si_oui),
    ],
    "STATUTS_CONSTITUTIFS_SARL_AU.docx": [
        ("anticipation_pluralite", transform_si_oui),
        ("apport_nature", transform_si_oui),
        ("fonds_commerce_apport", transform_si_oui),
        ("cac_nomme", transform_si_oui),
        ("acomptes_dividendes", transform_si_oui),
        ("tup_personne_morale", transform_si_oui),
        ("(si distincts de l'associé unique)", clean_inline_annotations),
        ("si distincts de l’associé unique", clean_inline_annotations),
        # Variantes inline → collapse en placeholder unique.
        ("apport_nature_evaluation", clean_inline_annotations),
        ("14.2 Désignation initiale", clean_inline_annotations),
    ],
    # JAL et Acte — aucune annotation NL identifiée à la passe initiale.
    # On scanne quand même pour sécurité (no-op si rien à faire).
    "ANNONCE_JAL_SARL.docx": [],
    "ANNONCE_JAL_SARL_AU.docx": [],
    "ACTE_NOMINATION_GERANT_SARL.docx": [],
    "ACTE_NOMINATION_GERANT_SARL_AU.docx": [],
}


def patch_template(path: Path, spec) -> dict:
    """Patche un .docx selon la spec. Retourne un dict de stats."""
    stats = {
        "file": path.name,
        "paragraphs_total": 0,
        "paragraphs_patched": 0,
        "matches": [],
    }
    if not path.exists():
        stats["error"] = f"introuvable : {path}"
        return stats

    doc = Document(path)

    def iter_paragraphs():
        for p in doc.paragraphs:
            yield p, "body"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p, "table"

    for para, origin in iter_paragraphs():
        stats["paragraphs_total"] += 1
        text = para_text(para)
        new_text = text
        applied = []
        for matcher, transform in spec:
            if matcher in text or matcher in new_text:
                if callable(transform):
                    candidate = transform(new_text)
                else:
                    candidate = transform
                if candidate != new_text:
                    applied.append(matcher)
                    new_text = candidate
        if new_text != text:
            rewrite_para_text(para, new_text)
            stats["paragraphs_patched"] += 1
            stats["matches"].append({
                "origin": origin,
                "applied": applied,
                "before": text[:200],
                "after": new_text[:200],
            })

    # Sauvegarde la version originale et écrit la nouvelle.
    backup = path.with_suffix(".preconv.docx")
    if not backup.exists():
        shutil.copy2(path, backup)
    doc.save(path)
    return stats


def main():
    if not TPL_DIR.exists():
        print(f"!! Dossier templates introuvable : {TPL_DIR}", file=sys.stderr)
        sys.exit(2)

    all_stats = []
    for fname, spec in SPEC.items():
        p = TPL_DIR / fname
        st = patch_template(p, spec)
        all_stats.append(st)
        print(f"== {fname}")
        if "error" in st:
            print(f"   ERR : {st['error']}")
            continue
        print(f"   {st['paragraphs_patched']}/{st['paragraphs_total']} paragraphes patchés")
        for m in st["matches"][:5]:
            print(f"   - matcher={m['applied']} origin={m['origin']}")
            print(f"     BEFORE: {m['before'][:120]}")
            print(f"     AFTER : {m['after'][:120]}")

    total = sum(s.get("paragraphs_patched", 0) for s in all_stats)
    print()
    print(f"Total paragraphes patchés : {total}")


if __name__ == "__main__":
    main()
